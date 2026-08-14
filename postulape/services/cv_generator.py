"""
Generador de CV adaptado a cada oferta.

Flujo:
  1. El LLM toma tu CV base + la oferta y produce un CV REESTRUCTURADO en JSON:
     reordena logros, prioriza los relevantes y usa las palabras clave de la
     oferta. NO inventa experiencia (guardrail explícito en el prompt).
  2. python-docx renderiza un .docx ATS-friendly: una sola columna, sin tablas
     ni cajas de texto, nombre y contacto en el cuerpo, viñetas simples.

    pip install python-docx
"""

import os
import re

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from postulape import config

SYSTEM = """Eres un experto en redacción de CVs optimizados para sistemas ATS.
Adaptas el CV de un candidato a una oferta específica.

REGLAS ESTRICTAS:
- NO inventes experiencia, títulos, empresas, fechas ni habilidades que no
  estén en el CV base. Solo puedes reordenar, reformular y resaltar lo REAL.
- Prioriza y reescribe los logros más relevantes para la oferta.
- Integra de forma natural las palabras clave de la oferta que el candidato
  realmente cumple (importante para pasar el ATS).
- Mantén un tono profesional y conciso, orientado a resultados.
- NO ELIMINES secciones que existan en el CV base: si el candidato tiene
  certificaciones, idiomas, logros o voluntariado, INCLÚYELOS (reordenados por
  relevancia). Son evidencia real y suelen ser decisivos.
- INCLUYE TODOS los empleos del CV base. Si uno es de otro rubro, mantenlo pero
  resume sus logros en 1 línea rescatando lo transferible (idiomas, atención al
  cliente, trabajo en equipo). No borres experiencia laboral.
- Las certificaciones y los idiomas se copian TAL CUAL del CV base (no los
  reformules ni los omitas): el ATS suele buscarlos literalmente."""

USER_TPL = """### CV BASE DEL CANDIDATO
{cv}

### OFERTA A LA QUE POSTULA
Título: {titulo}
Empresa: {empresa}
Descripción: {descripcion}
Palabras clave detectadas: {keywords}

Devuelve un JSON con EXACTAMENTE esta estructura (usa solo datos reales del CV base):
{{
  "nombre": "",
  "titular": "ej. Ingeniero de Software | Backend | Python",
  "contacto": {{"email": "", "telefono": "", "linkedin": "", "ubicacion": ""}},
  "resumen": "3-4 líneas adaptadas a esta oferta",
  "habilidades": ["skill1", "skill2", "..."],
  "experiencia": [
    {{"puesto": "", "empresa": "", "periodo": "",
      "logros": ["logro reescrito y orientado a la oferta", "..."]}}
  ],
  "educacion": [{{"titulo": "", "institucion": "", "periodo": ""}}],
  "proyectos": [{{"nombre": "", "descripcion": ""}}],
  "certificaciones": ["certificación tal cual figura en el CV base", "..."],
  "logros": ["logro/reconocimiento del CV base", "..."],
  "idiomas": ["ej. Español (nativo)", "ej. Inglés B2"],
  "voluntariado": [{{"rol": "", "organizacion": "", "periodo": "", "detalle": ""}}]
}}"""


def _slug(texto: str) -> str:
    texto = re.sub(r"[^\w\s-]", "", (texto or "").strip().lower())
    return re.sub(r"\s+", "_", texto)[:40] or "sin_nombre"


def _h(doc, texto):
    p = doc.add_paragraph()
    run = p.add_run(texto.upper())
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)
    p.space_before = Pt(10)
    p.space_after = Pt(2)
    # línea inferior como separador (regla, no tabla -> ATS-friendly)
    pPr = p._p.get_or_add_pPr()
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "999999")
    pbdr.append(bottom)
    pPr.append(pbdr)
    return p


def _render_docx(data: dict, ruta_salida: str):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    # Nombre
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(data.get("nombre", ""))
    run.bold = True
    run.font.size = Pt(20)

    # Titular
    if data.get("titular"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(data["titular"]).font.size = Pt(11)

    # Contacto (en el cuerpo, no en header -> ATS lo lee)
    c = data.get("contacto", {}) or {}
    linea = " | ".join(v for v in [c.get("email"), c.get("telefono"),
                                   c.get("linkedin"), c.get("ubicacion")] if v)
    if linea:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(linea).font.size = Pt(9.5)

    if data.get("resumen"):
        _h(doc, "Resumen profesional")
        doc.add_paragraph(data["resumen"])

    if data.get("habilidades"):
        _h(doc, "Habilidades")
        doc.add_paragraph(" · ".join(data["habilidades"]))

    if data.get("experiencia"):
        _h(doc, "Experiencia")
        for exp in data["experiencia"]:
            p = doc.add_paragraph()
            p.add_run(f"{exp.get('puesto','')} — {exp.get('empresa','')}").bold = True
            if exp.get("periodo"):
                p.add_run(f"  ({exp['periodo']})").italic = True
            for logro in exp.get("logros", []):
                doc.add_paragraph(logro, style="List Bullet")

    if data.get("proyectos"):
        _h(doc, "Proyectos")
        for pr in data["proyectos"]:
            p = doc.add_paragraph()
            p.add_run(pr.get("nombre", "")).bold = True
            if pr.get("descripcion"):
                p.add_run(f": {pr['descripcion']}")

    if data.get("educacion"):
        _h(doc, "Educación")
        for ed in data["educacion"]:
            p = doc.add_paragraph()
            p.add_run(ed.get("titulo", "")).bold = True
            resto = " — ".join(v for v in [ed.get("institucion"), ed.get("periodo")] if v)
            if resto:
                p.add_run(f" — {resto}")

    if data.get("certificaciones"):
        _h(doc, "Certificaciones")
        for cert in data["certificaciones"]:
            doc.add_paragraph(str(cert), style="List Bullet")

    if data.get("logros"):
        _h(doc, "Logros y reconocimientos")
        for lg in data["logros"]:
            doc.add_paragraph(str(lg), style="List Bullet")

    if data.get("idiomas"):
        _h(doc, "Idiomas")
        doc.add_paragraph(" · ".join(str(i) for i in data["idiomas"]))

    if data.get("voluntariado"):
        _h(doc, "Voluntariado")
        for v in data["voluntariado"]:
            p = doc.add_paragraph()
            p.add_run(f"{v.get('rol','')} — {v.get('organizacion','')}").bold = True
            if v.get("periodo"):
                p.add_run(f"  ({v['periodo']})").italic = True
            if v.get("detalle"):
                doc.add_paragraph(v["detalle"], style="List Bullet")

    os.makedirs(os.path.dirname(ruta_salida), exist_ok=True)
    doc.save(ruta_salida)


def generar(aviso: dict, cv_texto: str, llm, dir_cvs: str = None) -> str:
    """Genera el CV adaptado y devuelve la ruta del .docx (o '' si falla).
    dir_cvs: carpeta destino (por persona). Si es None usa config.DIR_CVS."""
    dir_cvs = dir_cvs or config.DIR_CVS
    user = USER_TPL.format(
        cv=cv_texto[:8000],
        titulo=aviso.get("titulo", ""),
        empresa=aviso.get("empresa", ""),
        descripcion=(aviso.get("descripcion") or "")[:4000],
        keywords=", ".join(aviso.get("keywords", [])),
    )
    try:
        data = llm.chat_json(SYSTEM, user)
    except Exception as e:
        print(f"[CV] Error generando CV para {aviso.get('id')}: {e}")
        return ""

    nombre = f"CV_{_slug(aviso.get('empresa'))}_{_slug(aviso.get('titulo'))}_{aviso.get('id')}.docx"
    ruta = os.path.join(dir_cvs, nombre)
    _render_docx(data, ruta)
    print(f"[CV] Generado: {ruta}")
    return ruta