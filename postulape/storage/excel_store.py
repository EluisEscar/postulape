"""
Almacenamiento en Excel (openpyxl). 11 columnas. Deduplica por 'id'.
Sanitiza texto para que Excel no interprete fórmulas (=, +, -, @).
Migra encabezados si el archivo existente tiene un esquema distinto.

    pip install openpyxl
"""

import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from postulape.status import es_estado_terminal

ENCABEZADOS = [
    "id", "plataforma", "titulo", "empresa", "distrito", "departamento",
    "modalidad", "antiguedad", "descripcion", "link", "estado",
]

_IDX_ID = ENCABEZADOS.index("id")
_MAX_CELDA = 32000
VERDE = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")


def _sanitizar(valor):
    """Evita inyección de fórmulas: si un texto empieza con = + - @, le antepone
    una comilla simple para que Excel lo trate como texto."""
    if isinstance(valor, str) and valor[:1] in ("=", "+", "-", "@"):
        return "'" + valor
    return valor


def _asegurar_libro(ruta: str):
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    if os.path.exists(ruta):
        wb = load_workbook(ruta)
        ws = wb["Postulaciones"] if "Postulaciones" in wb.sheetnames else wb.active
        ws.title = "Postulaciones"
        # Migración: si los encabezados actuales no coinciden, los reescribe.
        actuales = [c.value for c in ws[1]] if ws.max_row >= 1 else []
        if actuales[:len(ENCABEZADOS)] != ENCABEZADOS:
            for j, nombre in enumerate(ENCABEZADOS, start=1):
                ws.cell(row=1, column=j, value=nombre).font = Font(bold=True)
        return wb, ws
    wb = Workbook()
    ws = wb.active
    ws.title = "Postulaciones"
    ws.append(ENCABEZADOS)
    for celda in ws[1]:
        celda.font = Font(bold=True)
    ws.freeze_panes = "A2"
    wb.save(ruta)
    return wb, ws


def cargar_ids_existentes(ruta: str) -> set:
    """Lee la columna 'id' para saber qué avisos ya están guardados."""
    if not os.path.exists(ruta):
        return set()
    wb = load_workbook(ruta, read_only=True)
    ws = wb["Postulaciones"] if "Postulaciones" in wb.sheetnames else wb.active
    ids = set()
    for i, fila in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        if fila and len(fila) > _IDX_ID and fila[_IDX_ID]:
            ids.add(str(fila[_IDX_ID]))
    wb.close()
    return ids


def cargar_ids_procesados(ruta: str) -> set:
    """IDs con una decisión terminal; pendientes y errores son reintentables."""
    if not os.path.exists(ruta):
        return set()
    wb = load_workbook(ruta, read_only=True)
    ws = wb["Postulaciones"] if "Postulaciones" in wb.sheetnames else wb.active
    encabezados = [c.value for c in ws[1]]
    try:
        idx_id = encabezados.index("id")
        idx_estado = encabezados.index("estado")
    except ValueError:
        wb.close()
        return set()
    ids = {
        str(fila[idx_id])
        for fila in ws.iter_rows(min_row=2, values_only=True)
        if len(fila) > max(idx_id, idx_estado)
        and fila[idx_id]
        and es_estado_terminal(fila[idx_estado])
    }
    wb.close()
    return ids


def agregar_avisos(avisos: list, ruta: str, actualizar_existentes=False) -> dict:
    """Inserta avisos nuevos deduplicando por 'id'. Los que aplican se resaltan
    en verde. Guarda el veredicto resumido en 'estado'."""
    wb, ws = _asegurar_libro(ruta)
    filas_por_id = {
        str(ws.cell(row=i, column=_IDX_ID + 1).value): i
        for i in range(2, ws.max_row + 1)
        if ws.cell(row=i, column=_IDX_ID + 1).value
    }
    existentes = set(filas_por_id)

    insertados = actualizados = duplicados = 0
    for a in avisos:
        aviso_id = str(a.get("id") or "")
        if not aviso_id:
            duplicados += 1
            continue
        fila = [
            aviso_id,
            a.get("plataforma", ""),
            a.get("titulo", ""),
            a.get("empresa", ""),
            a.get("distrito", ""),
            a.get("departamento", ""),
            a.get("modalidad", ""),
            a.get("antiguedad", ""),
            (a.get("descripcion") or "")[:_MAX_CELDA],
            a.get("link", ""),
            a.get("estado", "Pendiente"),
        ]
        fila = [_sanitizar(v) for v in fila]
        if aviso_id in existentes:
            if not actualizar_existentes:
                duplicados += 1
                continue
            numero = filas_por_id[aviso_id]
            for columna, valor in enumerate(fila, start=1):
                ws.cell(row=numero, column=columna, value=valor)
            actualizados += 1
        else:
            ws.append(fila)
            filas_por_id[aviso_id] = ws.max_row
            existentes.add(aviso_id)
            insertados += 1
        if a.get("aplica"):
            numero = filas_por_id[aviso_id]
            for celda in ws[numero]:
                celda.fill = VERDE

    wb.save(ruta)
    print(f"[Excel] {insertados} insertados, {actualizados} actualizados, "
          f"{duplicados} duplicados omitidos.")
    return {"insertados": insertados, "actualizados": actualizados,
            "duplicados": duplicados}
