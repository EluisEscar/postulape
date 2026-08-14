"""
Prueba la subida de CVs a Supabase Storage (bucket PRIVADO) y el enlace temporal.
No scrapea ni usa el LLM.

Antes de correr:
  1. En Supabase → Storage → New bucket → nombre: cvs → PUBLIC DESACTIVADO.
  2. En el .env: SUPABASE_URL, SUPABASE_SERVICE_KEY, BUCKET_CVS=cvs

Uso:
    python -m postulape.test.test_storage                      # sube un .docx de prueba
    python -m postulape.test.test_storage output/cvs/x.docx    # sube un CV real
"""

import os
import sys

from postulape import config
from postulape.storage import db

def _crear_docx_prueba(ruta: str):
    from docx import Document
    doc = Document()
    doc.add_heading("CV de prueba", level=1)
    doc.add_paragraph("Este archivo es solo para probar Supabase Storage.")
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    doc.save(ruta)

def main():
    print("Bucket:", config.BUCKET_CVS, "| vigencia enlace:", config.CV_LINK_DIAS, "días")

    if len(sys.argv) > 1:
        ruta = sys.argv[1]
        if not os.path.exists(ruta):
            print(f"No existe el archivo: {ruta}")
            return
    else:
        ruta = os.path.join(config.DIR_CVS, "_prueba", "CV_prueba.docx")
        _crear_docx_prueba(ruta)
        print(f"Creado archivo de prueba: {ruta}")

    print("\nSubiendo a Storage...")
    ruta_storage = db.subir_cv(ruta, persona="prueba")
    if not ruta_storage:
        print("Falló la subida. Revisa que el bucket exista y las credenciales.")
        return

    print("\nGenerando enlace temporal firmado...")
    enlace = db.enlace_temporal(ruta_storage)
    if enlace:
        print("\n" + "=" * 70)
        print("ENLACE (caduca en", config.CV_LINK_DIAS, "días):")
        print(enlace)
        print("=" * 70)
        print("\n✔ Ábrelo en el navegador: debería descargar el .docx.")
        print("  Al caducar, el mismo enlace dejará de funcionar (bucket privado).")
    else:
        print("No se pudo generar el enlace.")

if __name__ == "__main__":
    main()